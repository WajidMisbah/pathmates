#!/usr/bin/env python3
"""Build data.json for the Pathmates PWA.
Usage: python3 build_data.py [--src DIR] [--progress FILE] [--out data.json]
Reads every '*.xlsx' in --src that has an 'All entities' sheet (cols: S.No, Category,
Subgroup, Entity, Variants, Remarks), merges with the organ queue in progress.json,
writes data.json and bumps the service-worker cache version (sw.js) so clients update."""
import argparse, glob, json, os, re, hashlib, datetime
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
CATS = ["Non-neoplastic - Inflammatory","Non-neoplastic - Others","Preneoplastic",
        "Neoplastic - Benign","Neoplastic - Malignant","Neoplastic - Others"]

def read_xlsx(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if "All entities" not in wb.sheetnames: return None
    rows = list(wb["All entities"].iter_rows(values_only=True))
    hdr = [str(h or "").strip().lower() for h in rows[0]]
    def col(name):
        for i,h in enumerate(hdr):
            if h.startswith(name): return i
        raise KeyError(name)
    ci = {k: col(k) for k in ["category","subgroup","entity","variants","remarks"]}
    for k,pre in [("histo","histopath"),("cyto","cytopath"),("ihc","immuno")]:
        i = next((j for j,h in enumerate(hdr) if h.startswith(pre)), None)
        if i is not None: ci[k] = i
    out = []
    for r in rows[1:]:
        if not r or not r[ci["entity"]]: continue
        out.append({k: (str(r[i]).strip() if r[i] is not None else "") for k,i in ci.items()})
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=os.path.join(HERE, "..", "organ-lists"))
    ap.add_argument("--progress", default=os.path.join(HERE, "..", "organ-lists", "progress.json"))
    ap.add_argument("--out", default=os.path.join(HERE, "data.json"))
    a = ap.parse_args()
    prog = json.load(open(a.progress))
    done = {c["index"]: c for c in prog.get("completed", [])}
    files = {}
    for p in sorted(glob.glob(os.path.join(a.src, "*.xlsx"))):
        m = re.match(r"(\d+)\s", os.path.basename(p))
        if m: files[int(m.group(1))] = p
    organs, entities = [], []
    for o in prog["organs"]:
        idx = o["index"]; ents = read_xlsx(files[idx]) if idx in files else None
        rec = {"index": idx, "name": o["name"], "covers": o.get("covers_from_Human_Organs_List",""),
               "status": "done" if ents else "pending", "count": len(ents or []),
               "link": done.get(idx, {}).get("link", "")}
        organs.append(rec)
        for n, e in enumerate(ents or []):
            e.update(id=f"{idx}-{n}", organ=o["name"], organIndex=idx); entities.append(e)
    data = {"generated": datetime.datetime.now().isoformat(timespec="seconds"),
            "categories": CATS, "organs": organs, "entities": entities}
    try:  # keep old timestamp if content unchanged, so no-op runs produce no diff
        old = json.load(open(a.out, encoding="utf-8"))
        if {k: old.get(k) for k in ("categories","organs","entities")} == {k: data[k] for k in ("categories","organs","entities")}:
            data["generated"] = old.get("generated", data["generated"])
    except Exception: pass
    s = json.dumps(data, ensure_ascii=False, separators=(",",":"))
    open(a.out, "w", encoding="utf-8").write(s)
    h = hashlib.sha1(s.encode())
    for f in ["index.html","style.css","app.js","manifest.json"]:
        h.update(open(os.path.join(HERE, f), "rb").read())
    ver = h.hexdigest()[:10]
    sw = os.path.join(HERE, "sw.js")
    if os.path.exists(sw):
        t = open(sw).read()
        t = re.sub(r'const VERSION = "[^"]*";', f'const VERSION = "{ver}";', t)
        open(sw, "w").write(t)
    print(f"{len(entities)} entities, {sum(o['status']=='done' for o in organs)}/{len(organs)} organs done; cache {ver}")
if __name__ == "__main__": main()
